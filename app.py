from flask import Flask, render_template, request, send_file
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook

app = Flask(__name__)


# 🔹 читаем только ВИДИМЫЕ строки (как в макросе)
def read_visible_excel(file):
    wb = load_workbook(file, data_only=True)
    ws = wb.active

    rows = []

    for row in ws.iter_rows():
        row_num = row[0].row

        # пропускаем скрытые строки
        if ws.row_dimensions[row_num].hidden:
            continue

        values = [cell.value for cell in row]
        rows.append(values)

    headers = rows[0]
    data = rows[1:]

    return pd.DataFrame(data, columns=headers)


def process_file(df):
    df.columns = [str(c).strip().lower().replace("\xa0", " ") for c in df.columns]

    col_art = "артикул поставщика"
    col_doc = "тип документа"
    col_qty = "кол-во"

    if col_art not in df.columns or col_doc not in df.columns or col_qty not in df.columns:
        raise ValueError("Не найдены нужные столбцы")

    sales = {}
    returns = {}

    for _, row in df.iterrows():
        art = str(row[col_art]).strip().replace("\xa0", "")
        doc = str(row[col_doc]).strip().lower().replace("\xa0", "")
        qty = row[col_qty]

        if art == "" or art.lower() == "nan":
            continue

        if pd.isna(qty):
            continue

        try:
            qty = float(qty)
        except:
            continue

        if doc == "продажа":
            sales[art] = sales.get(art, 0) + abs(qty)

        elif doc == "возврат":
            returns[art] = returns.get(art, 0) + abs(qty)

        elif doc == "" or doc == "nan":
            if qty > 0:
                sales[art] = sales.get(art, 0) + abs(qty)
            elif qty < 0:
                returns[art] = returns.get(art, 0) + abs(qty)

    all_keys = sorted(set(sales.keys()) | set(returns.keys()))

    result = []

    for art in all_keys:
        s = sales.get(art, 0)
        r = returns.get(art, 0)
        result.append([art, s, r, s - r])

    return pd.DataFrame(result, columns=[
        "Артикул поставщика",
        "Продажи",
        "Возвраты",
        "Продано (с учётом возвратов)"
    ])


def merge_files(df1, df2):
    all_keys = sorted(
        set(df1["Артикул поставщика"]) |
        set(df2["Артикул поставщика"])
    )

    result = []

    for art in all_keys:
        row1 = df1[df1["Артикул поставщика"] == art]
        row2 = df2[df2["Артикул поставщика"] == art]

        s1 = row1["Продажи"].sum() if not row1.empty else 0
        r1 = row1["Возвраты"].sum() if not row1.empty else 0
        n1 = row1["Продано (с учётом возвратов)"].sum() if not row1.empty else 0

        s2 = row2["Продажи"].sum() if not row2.empty else 0
        r2 = row2["Возвраты"].sum() if not row2.empty else 0
        n2 = row2["Продано (с учётом возвратов)"].sum() if not row2.empty else 0

        result.append([
            art,
            s1, r1, n1,
            s2, r2, n2,
            n1 + n2
        ])

    return pd.DataFrame(result, columns=[
        "Артикул поставщика",
        "Основной - Продажи",
        "Основной - Возвраты",
        "Основной Итог",
        "Доп - Продажи",
        "Доп - Возвраты",
        "Доп Итог",
        "Сумма Итог"
    ])


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/upload", methods=["POST"])
def upload():
    try:
        file1 = request.files["file1"]
        file2 = request.files["file2"]

        df1 = read_visible_excel(file1)
        df2 = read_visible_excel(file2)

        processed1 = process_file(df1)
        processed2 = process_file(df2)

        final_df = merge_files(processed1, processed2)

        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            final_df.to_excel(writer, index=False)

        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="Свод_Основной_и_Доп.xlsx"
        )

    except Exception as e:
        return f"Ошибка: {e}"


if __name__ == "__main__":
    app.run(debug=True)